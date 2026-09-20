"""Income Statement + Cash Summary for accountant-facing exports.

Uses the same ledger expansion as the Financials dashboard (`compute_financials`
via `finance_recurrence`). PDF rendering reuses the Weekly Pack ReportLab
pipeline. Excel is a real workbook (Income Statement, Cash Summary, Line items).

Deliberately not a balance sheet.
"""
from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from typing import Any, Optional

import finance_recurrence as fin_recur
from finance_entry import normalize_entry_name
from money_fmt import currency_symbol, entered_cash_amount, normalize_currency

import weekly_pack_export as pack_pdf

_FIN_FOOTER = (
    "Generated with Trenston. Income Statement and Cash Summary for the selected period. "
    "This is not a balance sheet."
)
_MONEY_FORMAT = '#,##0.00'
_HEADER_FILL = "8A7340"


def format_export_amount(n: float | None, currency: str = "usd") -> str:
    if n is None:
        return "—"
    sym = currency_symbol(currency)
    return f"{sym}{float(n):,.2f}"


def period_label(period: str) -> str:
    try:
        return datetime.strptime(period, "%Y-%m").strftime("%B %Y")
    except ValueError:
        return period


def financial_pdf_filename(workspace_name: str, period: str) -> str:
    slug = pack_pdf.slug_filename_part(workspace_name)
    return f"Trenston-Financial-Export-{slug}-{period}.pdf"


def financial_xlsx_filename(workspace_name: str, period: str) -> str:
    slug = pack_pdf.slug_filename_part(workspace_name)
    return f"Trenston-Financial-Export-{slug}-{period}.xlsx"


def _expand_ledger(entries: list[dict[str, Any]], now: Optional[datetime] = None):
    """Same expansion path as `compute_financials` (future months excluded from horizon)."""
    from collections import defaultdict

    valid = [e for e in (entries or []) if fin_recur.is_valid_month(str(e.get("month") or ""))]
    current, _scheduled = fin_recur.partition_ledger_entries(valid, now)
    horizon = fin_recur.resolve_expense_horizon(current, now)
    rev_by = defaultdict(float, fin_recur.expand_entries_by_month(current, entry_type="revenue", horizon_end=horizon))
    exp_by = defaultdict(float, fin_recur.expand_entries_by_month(current, entry_type="expense", horizon_end=horizon))
    cat_totals = fin_recur.expand_expense_category_totals(current, horizon)
    months = sorted(set(list(rev_by) + list(exp_by)))
    return current, horizon, months, rev_by, exp_by, cat_totals


def period_line_items(
    entries: list[dict[str, Any]],
    period: str,
    horizon: str,
) -> list[dict[str, Any]]:
    """Ledger rows that contribute to `period`, with name as the primary label.

    Uses the same non-overlapping recurring-rate supersession as dashboard totals
    so the Line items sheet cannot disagree with Income Statement aggregates.
    """
    raw = fin_recur.line_items_for_period(entries, period, horizon)
    rows: list[dict[str, Any]] = []
    for e in raw:
        category = (e.get("category") or "Other").strip() or "Other"
        rows.append({
            "name": normalize_entry_name(e.get("name"), category),
            "category": category,
            "type": e.get("type"),
            "amount": float(e.get("amount") or 0),
        })
    rows.sort(key=lambda r: (0 if r["type"] == "revenue" else 1, r["name"].lower(), r["category"].lower()))
    return rows


def reconstruct_cash_by_month(
    months: list[str],
    rev_by: dict[str, float],
    exp_by: dict[str, float],
    cash: float,
    period: str,
) -> tuple[Optional[float], Optional[float]]:
    """Walk cash backward from the dashboard snapshot (ending cash of the latest month).

    Returns (starting, ending) for `period`, or (None, None) when the period is
    after the latest ledger month (cannot project the snapshot forward).
    """
    if not months:
        return cash, cash
    latest = months[-1]
    if period > latest:
        return None, None
    first = min(months[0], period)
    walk = fin_recur.months_inclusive(first, latest)
    cursor_end = float(cash)
    starting = ending = None
    for m in reversed(walk):
        inflows = float(rev_by.get(m) or 0)
        outflows = float(exp_by.get(m) or 0)
        start = cursor_end - inflows + outflows
        if m == period:
            starting, ending = start, cursor_end
            break
        cursor_end = start
    return starting, ending


def assemble_financial_export(
    entries: list[dict[str, Any]],
    settings: dict | None,
    period: Optional[str] = None,
    *,
    now: Optional[datetime] = None,
) -> dict[str, Any]:
    settings = dict(settings or {})
    currency = normalize_currency(settings.get("currency"))
    _valid, horizon, months, rev_by, exp_by, cat_totals = _expand_ledger(entries, now)

    requested = (period or "").strip()
    if requested and not fin_recur.is_valid_month(requested):
        raise ValueError("Period must be a calendar month (YYYY-MM)")
    if not requested:
        requested = months[-1] if months else fin_recur.current_month(now)

    revenue = float(rev_by.get(requested) or 0)
    expenses_total = float(exp_by.get(requested) or 0)
    cats = cat_totals.get(requested) or {}
    by_category = [
        {"category": name, "amount": float(amt)}
        for name, amt in sorted(cats.items(), key=lambda kv: (-kv[1], kv[0]))
    ]
    # Keep category sum aligned with the month expense total used on Financials
    cat_sum = sum(row["amount"] for row in by_category)
    if by_category and abs(cat_sum - expenses_total) > 0.009:
        by_category.append({"category": "Other (unallocated)", "amount": round(expenses_total - cat_sum, 2)})

    net = revenue - expenses_total
    cash = entered_cash_amount(settings)
    inflows = revenue
    outflows = expenses_total
    start_cash = end_cash = None
    if cash is not None:
        start_cash, end_cash = reconstruct_cash_by_month(months, rev_by, exp_by, cash, requested)

    latest = months[-1] if months else None
    return {
        "period": requested,
        "period_label": period_label(requested),
        "currency": currency,
        "months": months,
        "latest_month": latest,
        "income": {
            "revenue": revenue,
            "expenses_by_category": by_category,
            "expenses_total": expenses_total,
            "net_income": net,
        },
        "line_items": period_line_items(_valid, requested, horizon),
        "cash": {
            "entered": cash is not None,
            "dashboard_cash": cash,
            "starting": start_cash,
            "inflows": inflows,
            "outflows": outflows,
            "ending": end_cash,
            "ending_matches_dashboard": (
                cash is not None and latest is not None and requested == latest and end_cash is not None
            ),
        },
    }


def statement_markdown(bundle: dict[str, Any]) -> str:
    currency = bundle["currency"]
    income = bundle["income"]
    cash = bundle["cash"]
    lines = [
        "# Income Statement",
        f"Period: **{bundle['period_label']}**",
        "",
        f"**Revenue**: {format_export_amount(income['revenue'], currency)}",
        "",
        "**Expenses**",
    ]
    if income["expenses_by_category"]:
        for row in income["expenses_by_category"]:
            lines.append(f"- {row['category']}: {format_export_amount(row['amount'], currency)}")
    else:
        lines.append("- None recorded this period")
    lines.extend(
        [
            "",
            f"**Total expenses**: {format_export_amount(income['expenses_total'], currency)}",
            f"**Net income**: {format_export_amount(income['net_income'], currency)}",
            "",
            "# Line items",
        ]
    )
    items = bundle.get("line_items") or []
    if items:
        for row in items:
            lines.append(
                f"- {row['name']} ({row['category']}, {row['type']}): "
                f"{format_export_amount(row['amount'], currency)}"
            )
    else:
        lines.append("- None recorded this period")
    lines.extend(
        [
            "",
            "# Cash Summary",
        ]
    )
    if cash["entered"]:
        if cash["starting"] is None and cash["ending"] is None:
            lines.append(
                "Starting and ending cash are shown only through the latest month on Financials; "
                "this period is after that snapshot."
            )
        else:
            lines.append(f"**Starting cash**: {format_export_amount(cash['starting'], currency)}")
        lines.append(f"**Inflows (revenue)**: {format_export_amount(cash['inflows'], currency)}")
        lines.append(f"**Outflows (expenses)**: {format_export_amount(cash['outflows'], currency)}")
        if cash["ending"] is not None:
            lines.append(f"**Ending cash**: {format_export_amount(cash['ending'], currency)}")
        if cash.get("ending_matches_dashboard"):
            lines.append("")
            lines.append(
                f"- Ending cash confirmed: matches Financials "
                f"({format_export_amount(cash['dashboard_cash'], currency)})"
            )
    else:
        lines.append("Cash on hand has not been entered on Financials, so starting and ending cash are omitted.")
        lines.append(f"**Inflows (revenue)**: {format_export_amount(cash['inflows'], currency)}")
        lines.append(f"**Outflows (expenses)**: {format_export_amount(cash['outflows'], currency)}")
        lines.append("**Starting cash**: —")
        lines.append("**Ending cash**: —")
    lines.extend(
        [
            "",
            "Scope is Income Statement and Cash Summary only. No balance sheet is included.",
        ]
    )
    return "\n".join(lines)


def render_financial_pdf(bundle: dict[str, Any], *, workspace_name: str) -> bytes:
    return pack_pdf.render_document_pdf(
        statement_markdown(bundle),
        workspace_name=workspace_name,
        kicker="Financial Export",
        pdf_title=f"Financial Export: {workspace_name} ({bundle['period']})",
        footer=_FIN_FOOTER,
        empty_message="Financial export is empty",
    )


def _xlsx_header_font():
    from openpyxl.styles import Font

    return Font(name="Calibri", bold=True, color="FFFFFF", size=12)


def _xlsx_title_font():
    from openpyxl.styles import Font

    return Font(name="Calibri", bold=True, size=16, color="18181B")


def _xlsx_label_font(*, bold=False):
    from openpyxl.styles import Font

    return Font(name="Calibri", bold=bold, size=11)


def _fill():
    from openpyxl.styles import PatternFill

    return PatternFill("solid", fgColor=_HEADER_FILL)


def _thin():
    from openpyxl.styles import Border, Side

    side = Side(style="thin", color="D4D4D8")
    return Border(left=side, right=side, top=side, bottom=side)


def _write_money(cell, value: float | None):
    if value is None:
        cell.value = "—"
        return
    cell.value = float(value)
    cell.number_format = _MONEY_FORMAT


def render_financial_xlsx(bundle: dict[str, Any], *, workspace_name: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    currency = bundle["currency"]
    sym = currency_symbol(currency)
    company = workspace_name or "Company"
    period = bundle["period_label"]
    income = bundle["income"]
    cash = bundle["cash"]
    border = _thin()
    fill = _fill()
    header_font = _xlsx_header_font()

    def style_header_row(ws, row: int, cols: int):
        for col in range(1, cols + 1):
            cell = ws.cell(row, col)
            cell.fill = fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.border = border

    def autosize(ws, cols: int):
        for col in range(1, cols + 1):
            letter = get_column_letter(col)
            widest = 12
            for cell in ws[letter]:
                val = cell.value
                widest = max(widest, min(48, len(str(val)) + 2 if val is not None else 0))
            ws.column_dimensions[letter].width = widest

    # --- Income Statement ---
    ws = wb.active
    ws.title = "Income Statement"
    ws["A1"] = company
    ws["A1"].font = _xlsx_title_font()
    ws["A2"] = f"Income Statement: {period}"
    ws["A2"].font = Font(name="Calibri", italic=True, size=11, color="52525B")
    ws["A3"] = f"Currency: {currency.upper()} ({sym})"
    ws["A3"].font = Font(name="Calibri", size=10, color="52525B")

    ws["A5"] = "Line"
    ws["B5"] = "Amount"
    style_header_row(ws, 5, 2)

    row = 6
    ws.cell(row, 1, "Revenue").font = _xlsx_label_font(bold=True)
    _write_money(ws.cell(row, 2), income["revenue"])
    ws.cell(row, 2).border = border
    ws.cell(row, 1).border = border
    row += 1
    ws.cell(row, 1, "Expenses by category").font = _xlsx_label_font(bold=True)
    ws.cell(row, 1).border = border
    ws.cell(row, 2).border = border
    row += 1
    if income["expenses_by_category"]:
        for item in income["expenses_by_category"]:
            ws.cell(row, 1, item["category"]).font = _xlsx_label_font()
            _write_money(ws.cell(row, 2), item["amount"])
            ws.cell(row, 1).border = border
            ws.cell(row, 2).border = border
            row += 1
    else:
        ws.cell(row, 1, "None recorded this period")
        ws.cell(row, 1).border = border
        ws.cell(row, 2).border = border
        row += 1
    ws.cell(row, 1, "Total expenses").font = _xlsx_label_font(bold=True)
    _write_money(ws.cell(row, 2), income["expenses_total"])
    ws.cell(row, 1).border = border
    ws.cell(row, 2).border = border
    row += 1
    ws.cell(row, 1, "Net income").font = _xlsx_label_font(bold=True)
    _write_money(ws.cell(row, 2), income["net_income"])
    ws.cell(row, 1).border = border
    ws.cell(row, 2).border = border
    row += 2
    ws.cell(row, 1, "Not a balance sheet. Income Statement only.")
    ws.cell(row, 1).font = Font(name="Calibri", italic=True, size=9, color="52525B")
    autosize(ws, 2)
    ws.freeze_panes = "A6"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1

    # --- Cash Summary ---
    cs = wb.create_sheet("Cash Summary")
    cs["A1"] = company
    cs["A1"].font = _xlsx_title_font()
    cs["A2"] = f"Cash Summary: {period}"
    cs["A2"].font = Font(name="Calibri", italic=True, size=11, color="52525B")
    cs["A3"] = f"Currency: {currency.upper()} ({sym})"
    cs["A3"].font = Font(name="Calibri", size=10, color="52525B")

    cs["A5"] = "Line"
    cs["B5"] = "Amount"
    style_header_row(cs, 5, 2)

    rows = [
        ("Starting cash", cash["starting"] if cash["entered"] else None),
        ("Inflows (revenue)", cash["inflows"]),
        ("Outflows (expenses)", cash["outflows"]),
        ("Ending cash", cash["ending"] if cash["entered"] else None),
    ]
    r = 6
    for label, amount in rows:
        cs.cell(r, 1, label).font = _xlsx_label_font(bold=label.startswith("Ending") or label.startswith("Starting"))
        _write_money(cs.cell(r, 2), amount)
        cs.cell(r, 1).border = border
        cs.cell(r, 2).border = border
        r += 1
    r += 1
    if not cash["entered"]:
        note = "Cash on hand has not been entered on Financials. Starting and ending cash are omitted so missing cash is not shown as $0."
    elif cash["ending_matches_dashboard"]:
        note = (
            f"Ending cash matches Cash on the Financials dashboard "
            f"({format_export_amount(cash['dashboard_cash'], currency)})."
        )
    elif cash["starting"] is None:
        note = "Starting and ending cash are reconstructed only through the latest month shown on Financials."
    else:
        note = (
            "Ending cash for the latest ledger month matches Cash on Financials; "
            "earlier months are walked backward from that snapshot using the same revenue and expenses."
        )
    cs.cell(r, 1, note)
    cs.cell(r, 1).font = Font(name="Calibri", italic=True, size=9, color="52525B")
    autosize(cs, 2)
    cs.freeze_panes = "A6"

    li = wb.create_sheet("Line items")
    li["A1"] = company
    li["A1"].font = _xlsx_title_font()
    li["A2"] = f"Line items: {period}"
    li["A2"].font = Font(name="Calibri", italic=True, size=11, color="52525B")
    li["A3"] = f"Currency: {currency.upper()} ({sym})"
    li["A3"].font = Font(name="Calibri", size=10, color="52525B")

    li["A5"] = "Name"
    li["B5"] = "Category"
    li["C5"] = "Type"
    li["D5"] = "Amount"
    style_header_row(li, 5, 4)
    items = bundle.get("line_items") or []
    r = 6
    if items:
        for item in items:
            li.cell(r, 1, item["name"]).font = _xlsx_label_font(bold=True)
            li.cell(r, 2, item["category"]).font = _xlsx_label_font()
            li.cell(r, 3, item["type"]).font = _xlsx_label_font()
            _write_money(li.cell(r, 4), item["amount"])
            for col in range(1, 5):
                li.cell(r, col).border = border
            r += 1
    else:
        li.cell(r, 1, "None recorded this period")
        li.cell(r, 1).border = border
    autosize(li, 4)
    li.freeze_panes = "A6"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
