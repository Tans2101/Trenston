"""Quick import (drop a file, get a company): reading, mapping, records, confirm and undo."""
import asyncio
import io
import os
from datetime import datetime

os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
os.environ.setdefault("DB_NAME", "test_quick_import")

import openpyxl  # noqa: E402
import pytest  # noqa: E402
from mongomock_motor import AsyncMongoMockClient  # noqa: E402
from starlette.datastructures import UploadFile  # noqa: E402

import quick_import as qi  # noqa: E402
import server  # noqa: E402


def _xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "P&L 2026"
    ws.append(["ACME Manufacturing - Profit and Loss"])
    ws.append([])
    ws.append(["Account", "Jan", "Feb", "Mar", "Nov", "Total"])
    ws.append(["Revenue"])
    ws.append(["Product sales", 120000, 135000, 128000, 150000, 533000])
    ws.append(["Total revenue", 120000, 135000, 128000, 150000, 533000])
    ws.append(["Salaries and wages", 60000, 60000, 62000, 62000, 0])
    ws.append(["Cost of sales", 30000, 32000, 31000, 33000, 0])
    ws.append(["Net income", 30000, 43000, 35000, 55000, 0])
    ws2 = wb.create_sheet("Clients")
    ws2.append(["Client", "Deal value (PHP)", "Status", "Sales rep", "Expected close"])
    ws2.append(["Jollibee Foods", "1,250,000", "Proposal sent", "Mark", datetime(2026, 11, 15)])
    ws2.append(["SM Supermalls", "800000", "Closed won", "Ana", "2026-09-30"])
    ws3 = wb.create_sheet("Job Orders")
    ws3.append(["JO #", "Product", "Qty", "Customer", "Delivery date"])
    ws3.append(["JO-1042", "Steel brackets", 5000, "Jollibee Foods", "2026-10-30"])
    ws4 = wb.create_sheet("Supplier orders")
    ws4.append(["Material", "Supplier", "Qty", "Total cost", "ETA"])
    ws4.append(["Cold rolled steel", "Pag-asa Steel", 20, "₱ 450,000.00", "2026-11-02"])
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def _analyze(raw: bytes, name: str):
    out = []
    for s in qi.read_workbook(raw, name):
        plan = qi.analyze_sheet(s, current_year=2026)
        built = qi.build_records(s, plan, current_month="2026-10", current_year=2026)
        out.append((s["name"], plan, built))
    return out


def test_reads_every_sheet_and_guesses_its_kind():
    kinds = {name: plan["kind"] for name, plan, _ in _analyze(_xlsx(), "acme.xlsx")}
    assert kinds == {
        "P&L 2026": qi.KIND_FINANCIAL_WIDE,
        "Clients": qi.KIND_DEALS,
        "Job Orders": qi.KIND_PRODUCTION,
        "Supplier orders": qi.KIND_PROCUREMENT,
    }


def test_pnl_skips_totals_and_future_months_and_adds_up():
    name, plan, built = _analyze(_xlsx(), "acme.xlsx")[0]
    recs = built["records"]
    assert all(r["month"] <= "2026-10" for r in recs)
    assert not any("Total" in r["name"] or "Net income" in r["name"] for r in recs)
    s = qi.summarize(plan["kind"], recs)
    assert s["revenue"] == 120000 + 135000 + 128000
    assert s["expenses"] == 60000 + 60000 + 62000 + 30000 + 32000 + 31000
    assert any("future months" in n for n in built["notes"])
    assert any("read as 2026" in n for n in built["notes"])
    reasons = " ".join(x["reason"] for x in built["skipped"])
    assert "Total revenue" in reasons and "Net income" in reasons


def test_deals_stage_value_and_dates():
    _, _, built = _analyze(_xlsx(), "acme.xlsx")[1]
    jolli, sm = built["records"]
    assert jolli["value"] == 1250000 and jolli["stage"] == "proposal" and jolli["close_date"] == "2026-11-15"
    assert sm["stage"] == "won" and sm["owner_name"] == "Ana"


def test_bank_statement_with_debit_and_credit_columns():
    csv = (b"Date,Description,Debit,Credit,Balance\n"
           b"2026-09-01,Opening balance,,,100000\n"
           b"2026-09-03,Meralco,12500.00,,87500\n"
           b"2026-09-05,Payment from Jollibee,,250000.00,337500\n"
           b"2026-12-01,Future,,5,0\n")
    (_, plan, built), = _analyze(csv, "bank.csv")
    assert plan["kind"] == qi.KIND_FINANCIAL
    types = [(r["type"], r["amount"]) for r in built["records"]]
    assert types == [("expense", 12500.0), ("revenue", 250000.0)]
    assert {s["reason"] for s in built["skipped"]} == {"No amount", "Date is in the future"}


def test_semicolon_task_list():
    csv = b"Task;Owner;Due;Status;Priority\nRenew permit;Tansher;2026-11-01;In progress;High\nCall supplier;Ana;;done;\n"
    (_, plan, built), = _analyze(csv, "todo.csv")
    assert plan["kind"] == qi.KIND_TASKS
    a, b = built["records"]
    assert a["column"] == "in_progress" and a["priority"] == "High" and a["due"] == "2026-11-01"
    assert b["column"] == "done"


def test_values_and_dates_are_parsed_carefully():
    assert qi.parse_amount("₱ 1,250.50") == 1250.5
    assert qi.parse_amount("(500)") == -500
    assert qi.parse_amount("abc") is None
    assert qi.parse_date("03/04/2026") is None  # ambiguous day and month
    assert str(qi.parse_date("25/12/2026")) == "2026-12-25"
    assert qi.parse_month("Jan", assumed_year=2026) == "2026-01"
    assert qi.parse_month("Sept 2025") == "2025-09"
    assert qi.classify_line("Cost of sales") == "expense"
    assert qi.classify_line("Gross profit") == "skip"
    assert qi.classify_line("Service income") == "revenue"


def test_rejects_old_excel_and_empty_files():
    with pytest.raises(ValueError, match="xlsx"):
        qi.read_workbook(b"\xd0\xcf\x11\xe0", "old.xls")
    with pytest.raises(ValueError):
        qi.read_workbook(b"", "x.csv")


def test_bad_ai_mapping_is_ignored():
    sheet = qi.read_workbook(b"Task,Due\nA,2026-01-01\n", "t.csv")[0]
    plan = qi.analyze_sheet(sheet, current_year=2026)
    # AI claims deals but maps no name column: keep the rules plan.
    merged = qi.merge_ai_plan(plan, {"kind": "deals", "columns": {"value": "Due"}}, current_year=2026)
    assert merged["kind"] == qi.KIND_TASKS
    # AI points at a header that does not exist: ignored too.
    merged = qi.merge_ai_plan(plan, {"kind": "tasks", "columns": {"title": "Nope"}}, current_year=2026)
    assert merged["kind"] == qi.KIND_TASKS and merged["source"] == "rules"


def test_clean_record_drops_invalid_and_future_rows():
    ok = qi.clean_record("financial", {"type": "revenue", "amount": "10", "month": "2026-09"}, current_month="2026-10")
    assert ok["amount"] == 10.0
    assert qi.clean_record("financial", {"type": "revenue", "amount": 10, "month": "2026-11"}, current_month="2026-10") is None
    assert qi.clean_record("financial", {"type": "gift", "amount": 10, "month": "2026-09"}, current_month="2026-10") is None
    assert qi.clean_record("deal", {"name": ""}, current_month="2026-10") is None
    assert qi.clean_record("deal", {"name": "X", "stage": "hacked", "value": -5}, current_month="2026-10")["stage"] == "lead"


def _setup(monkeypatch):
    db = AsyncMongoMockClient()["t"]
    monkeypatch.setattr(server, "db", db)

    async def fake_now(_ws):
        return datetime(2026, 10, 18, 9, 0)

    async def noop(*a, **k):
        return None

    monkeypatch.setattr(server, "_workspace_now", fake_now)
    monkeypatch.setattr(server, "log_activity", noop)
    monkeypatch.setattr(server, "_product_event", noop)
    monkeypatch.setattr(server.helm_llm, "anthropic_configured", lambda: False)
    return db


def test_preview_confirm_and_undo(monkeypatch):
    db = _setup(monkeypatch)
    owner = {"workspace_id": "w1", "user_id": "u1", "pack": "owner", "role": "owner", "name": "Tansher"}

    async def run():
        await db.workspaces.insert_one({"workspace_id": "w1", "name": "Acme", "tasks": {"items": [], "columns": []},
                                        "currency": "php"})
        prev = await server.quick_import_preview(
            UploadFile(file=io.BytesIO(_xlsx()), filename="acme.xlsx"), None, owner,
        )
        assert prev["used_ai"] is False
        sheets = prev["sheets"]
        assert [s["kind"] for s in sheets] == ["financial_wide", "deals", "production", "procurement"]
        assert sheets[2]["department_enabled"] is False
        body = server.QuickImportConfirmIn(
            filename="acme.xlsx",
            sheets=[server.QuickImportSheetIn(name=s["name"], kind=s["kind"], records=s["records"]) for s in sheets],
        )
        res = await server.quick_import_confirm(body, owner)
        batch = res["batch"]
        assert batch["counts"]["financial"] == 9 and batch["counts"]["deal"] == 2
        assert set(batch["enabled_departments"]) == {"production", "procurement"}
        assert await db.financial_entries.count_documents({"import_batch_id": batch["id"]}) == 9
        assert await db.production_work_orders.count_documents({"workspace_id": "w1"}) == 1
        # Same file again: no duplicates.
        again = await server.quick_import_confirm(body, owner)
        assert again["replayed"] is True
        assert await db.deals.count_documents({"workspace_id": "w1"}) == 2
        # Undo removes everything this import added.
        await server.quick_import_undo(batch["id"], owner)
        for coll in ("financial_entries", "deals", "production_work_orders", "procurement_requests"):
            assert await db[coll].count_documents({"workspace_id": "w1"}) == 0
        # After undo the same file can be imported again.
        res2 = await server.quick_import_confirm(body, owner)
        assert res2["replayed"] is False

    asyncio.run(run())


def test_only_the_owner_can_import(monkeypatch):
    _setup(monkeypatch)
    member = {"workspace_id": "w1", "user_id": "u2", "pack": "member", "role": "member"}

    async def run():
        with pytest.raises(server.HTTPException) as exc:
            await server.quick_import_preview(
                UploadFile(file=io.BytesIO(b"Task\nA\n"), filename="t.csv"), None, member,
            )
        assert exc.value.status_code == 403

    asyncio.run(run())


def test_ai_mapping_is_used_and_broken_ai_falls_back(monkeypatch):
    _setup(monkeypatch)
    monkeypatch.setattr(server.helm_llm, "anthropic_configured", lambda: True)
    owner = {"workspace_id": "w1", "user_id": "u1", "pack": "owner", "role": "owner"}
    csv = b"Customer,Amount due,Phase\nJollibee,5000,quoted\nSM,2000,signed\n"
    answers = iter([
        '{"kind": "deals", "columns": {"name": "Customer", "value": "Amount due", "stage": "Phase"}, "confidence": "high"}',
        "not json at all",
    ])

    async def fake_complete(system, user, **kw):
        assert "Jollibee" in user  # only a sample of rows is sent
        return next(answers)

    monkeypatch.setattr(server.helm_llm, "complete", fake_complete)

    async def run():
        prev = await server.quick_import_preview(UploadFile(file=io.BytesIO(csv), filename="c.csv"), None, owner)
        s = prev["sheets"][0]
        assert prev["used_ai"] is True and s["kind"] == "deals" and s["source"] == "ai"
        assert [r["stage"] for r in s["records"]] == ["proposal", "won"]
        prev2 = await server.quick_import_preview(UploadFile(file=io.BytesIO(csv), filename="c.csv"), None, owner)
        assert prev2["sheets"][0]["source"] == "rules"

    asyncio.run(run())
